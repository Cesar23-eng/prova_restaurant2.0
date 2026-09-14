import datetime
import os

from openpyxl import Workbook, load_workbook

from models import reports
from models.order import ORDER_TYPE_TAKEAWAY

DAY = datetime.date(2026, 9, 14)


def _sale(**overrides):
    sale = {
        "date": "2026-09-14", "time": "13:00:00", "number": "#0001", "table": "Mesa 1",
        "order_type": "En el local",
        "items": [{"dish": "Taco", "variant": "Carne", "note": "", "qty": 2,
                   "unit_price": 15.0, "subtotal": 30.0}],
        "total": 30.0, "method": "Efectivo", "cash_amount": 50.0, "qr_amount": 0.0,
        "amount_paid": 50.0, "change": 20.0, "change_method": "Efectivo",
        "moto_cost": None, "moto_payment_method": "",
    }
    sale.update(overrides)
    return sale


def test_cash_reconciliation_accounts_for_change_method():
    sales = [
        _sale(),  # 50 en efectivo, 20 de cambio en efectivo -> caja +30
        _sale(number="#0002", total=40.0, cash_amount=50.0, change=10.0,
              change_method="QR"),  # caja +50, QR -10
        _sale(number="#0003", total=25.0, method="QR", cash_amount=0.0,
              qr_amount=25.0, change=0.0, change_method=""),
        _sale(number="#0004", total=60.0, method="Mixto", cash_amount=30.0,
              qr_amount=40.0, change=10.0, change_method="Efectivo",
              order_type=ORDER_TYPE_TAKEAWAY, moto_cost=10.0, moto_payment_method="QR"),
    ]
    summary = reports.summarize(sales)
    assert summary["orders"] == 4
    assert summary["total"] == 155.0
    assert summary["cash_net"] == 30.0 + 50.0 + 0.0 + 20.0
    assert summary["qr_net"] == 0.0 - 10.0 + 25.0 + 40.0
    assert summary["cash_net"] + summary["qr_net"] == summary["total"]
    assert summary["moto_qr"] == 10.0 and summary["moto_orders"] == 1
    assert summary["by_method"]["Mixto"] == {"orders": 1, "total": 60.0}
    assert summary["by_type"][ORDER_TYPE_TAKEAWAY]["total"] == 60.0
    assert summary["products"][0] == {"product": "Taco (Carne)", "qty": 8, "amount": 120.0}


def test_daily_excel_is_rebuilt_from_ledger(tmp_path):
    root = str(tmp_path)
    reports.append_sale(root, DAY, _sale(items=[{
        "dish": "Taco", "variant": "Carne", "note": "sin cebolla", "qty": 2,
        "unit_price": 15.0, "subtotal": 30.0}]))
    reports.append_sale(root, DAY, _sale(number="#0002", order_type=ORDER_TYPE_TAKEAWAY,
                                         moto_cost=12.0, moto_payment_method="Efectivo"))
    ok, message = reports.write_daily_excel(root, DAY)
    assert ok, message
    wb = load_workbook(reports.excel_path(root, DAY))
    assert wb.sheetnames == ["En el local", "Para llevar", "Resumen"]
    local = list(wb["En el local"].iter_rows(values_only=True))
    assert local[1][2] == "Taco (Carne) x2 [sin cebolla]"
    takeaway = list(wb["Para llevar"].iter_rows(values_only=True))
    assert takeaway[1][9] == 12.0


def test_locked_excel_reports_warning_but_keeps_sale(tmp_path, monkeypatch):
    root = str(tmp_path)
    reports.append_sale(root, DAY, _sale())

    def locked(src, dst):
        raise PermissionError("abierto en Excel")

    monkeypatch.setattr(reports.os, "replace", locked)
    ok, message = reports.write_daily_excel(root, DAY)
    assert not ok and "abierto" in message
    assert len(reports.read_sales(root, DAY)) == 1


def test_legacy_excel_without_ledger_is_preserved(tmp_path):
    root = str(tmp_path)
    legacy = reports.excel_path(root, DAY)
    wb = Workbook()
    wb.active.append(["venta de la version anterior"])
    wb.save(legacy)
    reports.append_sale(root, DAY, _sale())
    folder = os.path.dirname(legacy)
    assert os.path.exists(os.path.join(folder, "pedidos_2026-09-14_anterior.xlsx"))


def test_truncated_ledger_line_is_skipped(tmp_path):
    root = str(tmp_path)
    reports.append_sale(root, DAY, _sale())
    with open(reports.ledger_path(root, DAY), "a", encoding="utf-8") as f:
        f.write('{"number": "#0002", "tot')
    assert len(reports.read_sales(root, DAY)) == 1
    assert reports.max_order_number(root, DAY) == 1
