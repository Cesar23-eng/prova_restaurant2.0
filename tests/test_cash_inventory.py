import datetime
import json
import os

import pytest
from openpyxl import load_workbook

from models.cash import cash_summary
from models.inventory import UnavailableError
from models.order import OrderManager
from utils import tickets

NOW = datetime.datetime(2026, 9, 23, 22, 0)


def sell(manager, table, qty=2, method="Efectivo", **payment):
    manager.create_table(table)
    manager.add_item(table, "Platillos", "Taco", "Pastor", 15, qty=qty)
    return manager.register_payment(table, method, **payment)


# ---------------------------------------------------------------------------
#  Caja
# ---------------------------------------------------------------------------
def test_open_cash_register_and_expected_amounts(manager):
    assert not manager.cash.is_open()
    manager.cash.open(cash=200, qr=50, reserve=100, cashier="Ana")
    sell(manager, "Mesa 1", 4, cash_amount=100, change_method="Efectivo")  # 60 de venta, 40 de cambio
    sell(manager, "Mesa 2", 2, method="QR", qr_amount=30)
    manager.cash.add_expense(15, "Efectivo", "Hielo")
    manager.cash.add_expense(20, "QR", "Gas")
    summary = manager.cash_summary()
    assert summary["opened"] and summary["cashier"] == "Ana"
    assert summary["expected_cash"] == 200 + 100 + 60 - 15
    assert summary["expected_qr"] == 50 + 30 - 20
    assert summary["reserve_left"] == 100 - 35
    assert summary["counted_cash"] is None and summary["difference"] is None


def test_motos_paid_from_the_drawer_are_discounted(manager):
    manager.cash.open(cash=100)
    manager.create_table("Juan", order_type="Para llevar")
    manager.add_item("Juan", "Platillos", "Taco", "Pastor", 15, qty=2)
    manager.set_delivery_details("Juan", 10, "Efectivo")
    manager.register_payment("Juan", "Efectivo", cash_amount=30)
    manager.create_table("Luis", order_type="Para llevar")
    manager.add_item("Luis", "Platillos", "Taco", "Pastor", 15, qty=2)
    manager.set_delivery_details("Luis", 8, "QR")
    manager.register_payment("Luis", "QR", qr_amount=30)
    summary = manager.cash_summary()
    assert summary["expected_cash"] == 100 + 30 - 10
    assert summary["expected_qr"] == 30 - 8
    off = cash_summary(manager.cash.load(), manager.day_summary(), motos_from_drawer=False)
    assert off["expected_cash"] == 130 and off["expected_qr"] == 30


def test_count_shows_shortage_or_surplus(manager):
    manager.cash.open(cash=100)
    sell(manager, "Mesa 1", 2, cash_amount=30)
    manager.cash.set_count(125)
    summary = manager.cash_summary()
    assert summary["expected_cash"] == 130 and summary["difference"] == -5
    manager.cash.set_count(131.5)
    assert manager.cash_summary()["difference"] == 1.5


def test_void_expense_and_validations(manager):
    expense = manager.cash.add_expense("12,50", "Efectivo", "  Servilletas  ")
    assert expense["amount"] == 12.5 and expense["reason"] == "Servilletas"
    assert manager.cash.void_expense(expense["id"])
    assert not manager.cash.void_expense(expense["id"])
    assert manager.cash_summary()["expenses_cash"] == 0
    for amount, reason in ((0, "x"), (5, ""), (-3, "x")):
        with pytest.raises(ValueError):
            manager.cash.add_expense(amount, "Efectivo", reason)
    with pytest.raises(ValueError):
        manager.cash.add_expense(5, "Tarjeta", "x")


def test_correcting_the_opening_keeps_original_time_and_survives_restart(manager, data_dir, frozen_now):
    manager.cash.open(cash=100)
    first = manager.cash.load()["opening"]["time"]
    manager.cash.open(cash=150, reserve=50)
    opening = manager.cash.load()["opening"]
    assert opening["cash"] == 150 and opening["time"] == first and "corrected_at" in opening
    restarted = OrderManager(root=data_dir, cutoff_hour=4)
    assert restarted.cash.is_open() and restarted.cash_summary()["expected_cash"] == 200
    day = manager.today().isoformat()
    with open(os.path.join(data_dir, day, f"audit_log_{day}.txt"), encoding="utf-8") as f:
        audit = f.read()
    assert "CAJA_ABRE" in audit and "CAJA_CORRIGE" in audit


def test_each_business_day_has_its_own_cash_register(manager, frozen_now):
    manager.cash.open(cash=100)
    frozen_now(datetime.datetime(2026, 9, 15, 12, 0))
    assert not manager.cash.is_open()
    assert manager.cash.is_open(datetime.date(2026, 9, 14))


def test_without_opening_the_close_is_the_same_as_before(manager):
    sell(manager, "Mesa 1", 2, cash_amount=50, change_method="Efectivo")
    summary = manager.cash_summary()
    assert not summary["opened"] and summary["expected_cash"] == manager.day_summary()["cash_net"]


def test_excel_and_ticket_include_cash_block(manager):
    manager.cash.open(cash=200, reserve=50, cashier="Ana")
    sell(manager, "Mesa 1", 2, cash_amount=30)
    manager.cash.add_expense(15, "Efectivo", "Hielo")
    manager.cash.set_count(260)
    ok, message = manager.refresh_daily_excel()
    assert ok, message
    rows = list(load_workbook(manager.daily_excel_path())["Resumen"].iter_rows(values_only=True))
    values = {row[0]: row[1] for row in rows if row and row[0]}
    assert values["Efectivo esperado"] == 265 and values["Diferencia (+ sobra / - falta)"] == -5
    assert any(row[0] == "Hielo" for row in rows if row)

    text = tickets.to_text(tickets.day_summary_ticket(
        "PRÖVA", manager.day_summary(), "23/09/2026", 48, NOW, cash=manager.cash_summary()))
    assert all(len(row) <= 48 for row in text.splitlines())
    assert "EFECTIVO ESPERADO" in text and "Bs 265.00" in text
    assert "FALTANTE" in text and "Bs 5.00" in text and "Hielo" in text


def test_ticket_prints_cash_block_even_without_sales(manager):
    manager.cash.open(cash=100)
    text = tickets.to_text(tickets.day_summary_ticket(
        "PRÖVA", manager.day_summary(), "23/09/2026", 48, NOW, cash=manager.cash_summary()))
    assert "EFECTIVO ESPERADO" in text and "No hay ventas cobradas" in text


# ---------------------------------------------------------------------------
#  Inventario (agotado / disponible)
# ---------------------------------------------------------------------------
def test_unavailable_soda_cannot_be_ordered(manager):
    manager.create_table("Mesa 1")
    manager.inventory.set_available("Bebidas", "Coca cola", "Botella", False)
    with pytest.raises(UnavailableError) as error:
        manager.add_items("Mesa 1", [
            {"category": "Platillos", "dish": "Taco", "variant": "Pastor", "price": 15},
            {"category": "Bebidas", "dish": "Coca cola", "variant": "Botella", "price": 10},
        ])
    assert error.value.items == [("Bebidas", "Coca cola", "Botella")]
    assert "Se acabó: Coca cola (Botella)" in str(error.value)
    assert manager.get_items("Mesa 1") == []  # el lote completo se rechaza
    manager.inventory.set_available("bebidas", " coca  COLA", "botella", True)
    manager.add_item("Mesa 1", "Bebidas", "Coca cola", "Botella", 10)
    assert len(manager.get_items("Mesa 1")) == 1


def test_unavailable_blocks_plus_button_on_existing_line(manager):
    manager.create_table("Mesa 1")
    manager.add_item("Mesa 1", "Bebidas", "Coca cola", "Botella", 10)
    manager.inventory.set_available("Bebidas", "Coca cola", "Botella", False)
    with pytest.raises(UnavailableError):
        manager.add_to_line("Mesa 1", 0)
    assert manager.remove_line("Mesa 1", 0) == 1  # quitar si se puede


def test_inventory_persists_across_days_and_restarts(manager, data_dir, frozen_now):
    events = []
    manager.add_listener(lambda event, table, info: events.append(event))
    assert manager.inventory.set_available("Bebidas", "Fanta papaya", "Botella", False)
    assert not manager.inventory.set_available("Bebidas", "Fanta papaya", "Botella", False)
    assert events == ["inventory"]
    frozen_now(datetime.datetime(2026, 9, 20, 12, 0))
    restarted = OrderManager(root=data_dir, cutoff_hour=4)
    assert not restarted.inventory.is_available("Bebidas", "Fanta papaya", "Botella")
    assert restarted.inventory.unavailable()[0]["platillo"] == "Fanta papaya"
    assert restarted.inventory.mark_all_available() == 1
    with open(os.path.join(data_dir, "inventario.json"), encoding="utf-8") as f:
        assert json.load(f) == {"agotados": []}
