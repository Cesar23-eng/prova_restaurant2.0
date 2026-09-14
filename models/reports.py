# models/reports.py
"""
Registro de ventas del dia y reportes.

Cada pago se agrega a data/AAAA-MM-DD/ventas_AAAA-MM-DD.jsonl (una linea JSON
por venta). Ese archivo es la fuente de verdad: el Excel del dia se regenera
completo a partir de el. Si el Excel esta abierto y Windows no deja
guardarlo, la venta igual queda registrada y el Excel se pone al dia en el
siguiente guardado.
"""
import datetime
import json
import os
import time
from collections import OrderedDict
from typing import Dict, List, Tuple

from utils.config import day_dir

SHEET_LOCAL = "En el local"
SHEET_TAKEAWAY = "Para llevar"
SHEET_SUMMARY = "Resumen"

HEADERS_LOCAL = [
    "#", "Mesa/Cliente", "Platillos", "Total (Bs)",
    "Metodo de Pago", "Efectivo recibido (Bs)", "Monto QR (Bs)",
    "Cambio (Bs)", "Cambio dado en", "Pagado", "Hora",
]
HEADERS_TAKEAWAY = [
    "#", "Mesa/Cliente", "Platillos", "Total (Bs)",
    "Metodo de Pago", "Efectivo recibido (Bs)", "Monto QR (Bs)",
    "Cambio (Bs)", "Cambio dado en",
    "Costo Moto (Bs)", "Pago Moto (metodo)", "Pagado", "Hora",
]


def _date_str(date: datetime.date) -> str:
    return date.strftime("%Y-%m-%d")


def ledger_path(root: str, date: datetime.date) -> str:
    return os.path.join(day_dir(root, date), f"ventas_{_date_str(date)}.jsonl")


def excel_path(root: str, date: datetime.date) -> str:
    return os.path.join(day_dir(root, date), f"pedidos_{_date_str(date)}.xlsx")


# ---------------------------------------------------------------------------
#  Registro (ledger)
# ---------------------------------------------------------------------------
def append_sale(root: str, date: datetime.date, sale: Dict, retries: int = 5):
    path = ledger_path(root, date)
    if not os.path.exists(path):
        _preserve_legacy_excel(root, date)
    line = json.dumps(sale, ensure_ascii=False) + "\n"
    for attempt in range(retries):
        try:
            with open(path, "a", encoding="utf-8") as f:
                f.write(line)
                f.flush()
                os.fsync(f.fileno())
            return
        except PermissionError:
            if attempt == retries - 1:
                raise
            time.sleep(0.1 * (attempt + 1))


def _preserve_legacy_excel(root: str, date: datetime.date):
    """
    Un Excel del dia sin registro .jsonl lo genero una version anterior de la
    app. Se renombra en vez de sobrescribirlo para no perder esas ventas.
    """
    path = excel_path(root, date)
    if not os.path.exists(path):
        return
    base, ext = os.path.splitext(path)
    target = f"{base}_anterior{ext}"
    i = 2
    while os.path.exists(target):
        target = f"{base}_anterior_{i}{ext}"
        i += 1
    try:
        os.replace(path, target)
    except OSError:
        pass


def read_sales(root: str, date: datetime.date) -> List[Dict]:
    path = ledger_path(root, date)
    if not os.path.exists(path):
        return []
    sales = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                sales.append(json.loads(line))
            except ValueError:
                # Una linea cortada por un apagon no invalida el resto del dia
                continue
    return sales


def max_order_number(root: str, date: datetime.date) -> int:
    best = 0
    for sale in read_sales(root, date):
        best = max(best, parse_order_number(sale.get("number", "")))
    return best


def parse_order_number(number: str) -> int:
    digits = "".join(ch for ch in str(number) if ch.isdigit())
    return int(digits) if digits else 0


# ---------------------------------------------------------------------------
#  Resumen / cierre de caja
# ---------------------------------------------------------------------------
def items_text(lines: List[Dict]) -> str:
    # En el Excel no importa en que plato se sirvio: se suman las mismas lineas
    merged: "OrderedDict[Tuple[str, str, str], int]" = OrderedDict()
    for line in lines:
        key = (line["dish"], line["variant"], line.get("note", ""))
        merged[key] = merged.get(key, 0) + int(line["qty"])
    parts = []
    for (dish, variant, note), qty in merged.items():
        text = f"{dish} ({variant}) x{qty}"
        if note:
            text += f" [{note}]"
        parts.append(text)
    return "; ".join(parts)


def summarize(sales: List[Dict]) -> Dict:
    by_type = OrderedDict((t, {"orders": 0, "total": 0.0}) for t in (SHEET_LOCAL, SHEET_TAKEAWAY))
    by_method = OrderedDict((m, {"orders": 0, "total": 0.0}) for m in ("Efectivo", "QR", "Mixto"))
    products: Dict[str, Dict] = {}
    summary = {
        "orders": len(sales),
        "total": 0.0,
        "average_ticket": 0.0,
        "cash_net": 0.0,
        "qr_net": 0.0,
        "change_cash": 0.0,
        "change_qr": 0.0,
        "moto_cash": 0.0,
        "moto_qr": 0.0,
        "moto_orders": 0,
        "by_type": by_type,
        "by_method": by_method,
        "products": [],
        "first_time": "",
        "last_time": "",
    }
    times = []
    for sale in sales:
        total = float(sale.get("total") or 0)
        change = float(sale.get("change") or 0)
        change_method = sale.get("change_method", "")
        summary["total"] += total

        order_type = sale.get("order_type") or SHEET_LOCAL
        bucket = by_type.setdefault(order_type, {"orders": 0, "total": 0.0})
        bucket["orders"] += 1
        bucket["total"] += total

        method = sale.get("method") or "Otro"
        bucket = by_method.setdefault(method, {"orders": 0, "total": 0.0})
        bucket["orders"] += 1
        bucket["total"] += total

        # Lo que realmente entro a caja y a la cuenta QR, descontando el cambio
        cash_in = float(sale.get("cash_amount") or 0)
        qr_in = float(sale.get("qr_amount") or 0)
        if change_method == "Efectivo":
            cash_in -= change
            summary["change_cash"] += change
        elif change_method == "QR":
            qr_in -= change
            summary["change_qr"] += change
        summary["cash_net"] += cash_in
        summary["qr_net"] += qr_in

        moto_cost = sale.get("moto_cost")
        if moto_cost not in (None, ""):
            summary["moto_orders"] += 1
            if sale.get("moto_payment_method") == "QR":
                summary["moto_qr"] += float(moto_cost)
            else:
                summary["moto_cash"] += float(moto_cost)

        for line in sale.get("items", []):
            key = f"{line['dish']} ({line['variant']})"
            product = products.setdefault(key, {"product": key, "qty": 0, "amount": 0.0})
            product["qty"] += int(line.get("qty", 0))
            product["amount"] += float(line.get("subtotal", 0))

        if sale.get("time"):
            times.append(sale["time"])

    if sales:
        summary["average_ticket"] = summary["total"] / len(sales)
    if times:
        summary["first_time"], summary["last_time"] = min(times), max(times)
    summary["products"] = sorted(products.values(), key=lambda p: (-p["qty"], -p["amount"]))
    for key in ("total", "average_ticket", "cash_net", "qr_net", "change_cash",
                "change_qr", "moto_cash", "moto_qr"):
        summary[key] = round(summary[key], 2)
    return summary


# ---------------------------------------------------------------------------
#  Excel
# ---------------------------------------------------------------------------
def _style_header(ws, headers):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A5276")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws):
    from openpyxl.utils import get_column_letter

    for col in ws.columns:
        max_len = max((len(str(c.value if c.value is not None else "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def _money(value):
    return round(float(value), 2) if value not in (None, "") else ""


def _nonzero(value):
    return _money(value) if value else ""


def _order_row(record: Dict, paid_label: str, with_time: bool) -> Tuple[str, list]:
    is_takeaway = record.get("order_type") == SHEET_TAKEAWAY
    row = [
        record.get("number", ""),
        record.get("table", ""),
        items_text(record.get("items", [])),
        _money(record.get("total", 0)),
        record.get("method", ""),
        _nonzero(record.get("cash_amount")),
        _nonzero(record.get("qr_amount")),
        _nonzero(record.get("change")),
        record.get("change_method", ""),
    ]
    if is_takeaway:
        row += [_money(record.get("moto_cost")), record.get("moto_payment_method", "")]
    row.append(paid_label)
    if with_time:
        row.append(record.get("time", ""))
    return (SHEET_TAKEAWAY if is_takeaway else SHEET_LOCAL), row


def _write_summary_sheet(ws, summary: Dict, date_label: str):
    from openpyxl.styles import Font

    bold = Font(bold=True)
    ws.append([f"Cierre de caja - {date_label}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    rows = [
        ("Pedidos cobrados", summary["orders"]),
        ("Total vendido (Bs)", summary["total"]),
        ("Ticket promedio (Bs)", summary["average_ticket"]),
        ("Primera venta", summary["first_time"]),
        ("Ultima venta", summary["last_time"]),
        (None, None),
        ("Efectivo neto en caja (Bs)", summary["cash_net"]),
        ("QR neto (Bs)", summary["qr_net"]),
        ("Cambio dado en efectivo (Bs)", summary["change_cash"]),
        ("Cambio dado por QR (Bs)", summary["change_qr"]),
        (None, None),
        ("Motos pagadas en efectivo (Bs)", summary["moto_cash"]),
        ("Motos pagadas por QR (Bs)", summary["moto_qr"]),
    ]
    for label, value in rows:
        ws.append([label, value] if label else [])
    ws.append([])
    ws.append(["Por metodo de pago", "Pedidos", "Total (Bs)"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for method, data in summary["by_method"].items():
        ws.append([method, data["orders"], round(data["total"], 2)])
    ws.append([])
    ws.append(["Por tipo de consumo", "Pedidos", "Total (Bs)"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for order_type, data in summary["by_type"].items():
        ws.append([order_type, data["orders"], round(data["total"], 2)])
    ws.append([])
    ws.append(["Producto", "Cantidad", "Importe (Bs)"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for product in summary["products"]:
        ws.append([product["product"], product["qty"], round(product["amount"], 2)])
    _autosize(ws)


def _save_workbook(wb, path: str) -> Tuple[bool, str]:
    tmp = f"{path}.tmp.xlsx"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
        return True, ""
    except PermissionError:
        try:
            os.remove(tmp)
        except OSError:
            pass
        return False, (
            f"No se pudo actualizar {os.path.basename(path)} porque esta abierto "
            f"en otro programa. Cierralo: la venta ya quedo registrada y el Excel "
            f"se actualizara en el proximo guardado."
        )
    except OSError as e:
        return False, f"No se pudo guardar el Excel: {e}"


def write_daily_excel(root: str, date: datetime.date) -> Tuple[bool, str]:
    from openpyxl import Workbook

    sales = read_sales(root, date)
    if not sales and os.path.exists(excel_path(root, date)):
        # Excel de una version anterior sin registro: no se pisa con uno vacio
        return True, ""
    wb = Workbook()
    ws_local = wb.active
    ws_local.title = SHEET_LOCAL
    ws_takeaway = wb.create_sheet(SHEET_TAKEAWAY)
    _style_header(ws_local, HEADERS_LOCAL)
    _style_header(ws_takeaway, HEADERS_TAKEAWAY)
    sheets = {SHEET_LOCAL: ws_local, SHEET_TAKEAWAY: ws_takeaway}
    for sale in sales:
        sheet, row = _order_row(sale, "Si", with_time=True)
        sheets[sheet].append(row)
    _autosize(ws_local)
    _autosize(ws_takeaway)
    _write_summary_sheet(wb.create_sheet(SHEET_SUMMARY), summarize(sales), _date_str(date))
    return _save_workbook(wb, excel_path(root, date))


def write_orders_snapshot(path: str, orders: List[Dict]) -> Tuple[bool, str]:
    """Respaldo manual con todos los pedidos en pantalla, pagados o no."""
    from openpyxl import Workbook

    wb = Workbook()
    ws_local = wb.active
    ws_local.title = SHEET_LOCAL
    ws_takeaway = wb.create_sheet(SHEET_TAKEAWAY)
    _style_header(ws_local, HEADERS_LOCAL)
    _style_header(ws_takeaway, HEADERS_TAKEAWAY)
    sheets = {SHEET_LOCAL: ws_local, SHEET_TAKEAWAY: ws_takeaway}
    for order in orders:
        if not order.get("items"):
            continue
        sheet, row = _order_row(order, "Si" if order.get("paid") else "No", with_time=True)
        sheets[sheet].append(row)
    _autosize(ws_local)
    _autosize(ws_takeaway)
    return _save_workbook(wb, path)
