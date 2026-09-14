import os
import threading
from typing import Dict, Optional

from utils.config import app_dir

MENU_FILENAME = "menu_precios.xlsx"
DEFAULT_VARIANT = "Normal"

Menu = Dict[str, Dict[str, Dict[str, float]]]


class MenuData:
    """
    Lee menu_precios.xlsx (columnas: categoria, producto, variante, precio).

    El menu queda en cache y solo se vuelve a leer si el archivo cambia, asi
    el dueño puede actualizar precios con la app abierta sin reiniciarla.
    """

    def __init__(self, path: Optional[str] = None):
        self._explicit_path = path
        self._lock = threading.Lock()
        self._menu: Menu = {}
        self._mtime: Optional[float] = None
        self.last_error = ""

    @staticmethod
    def resource_path(relative_path: str) -> str:
        """Busca junto al ejecutable; en desarrollo tambien en models/."""
        candidates = [
            os.path.join(app_dir(), relative_path),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), relative_path),
        ]
        for candidate in candidates:
            if os.path.exists(candidate):
                return candidate
        return candidates[0]

    @property
    def path(self) -> str:
        return self._explicit_path or self.resource_path(MENU_FILENAME)

    def get_menu_prices(self) -> Menu:
        path = self.path
        with self._lock:
            try:
                mtime = os.path.getmtime(path)
            except OSError:
                self.last_error = f"No se encontro el menu: {path}"
                return self._menu
            if mtime != self._mtime:
                try:
                    self._menu = self._read(path)
                    self._mtime = mtime
                    self.last_error = ""
                except Exception as e:
                    # Si el Excel esta a medio guardar se conserva el menu anterior
                    self.last_error = f"Error leyendo el menu: {e}"
            return self._menu

    def get_price(self, category: str, dish: str, variant: str) -> Optional[float]:
        try:
            return self.get_menu_prices()[category][dish][variant]
        except KeyError:
            return None

    @staticmethod
    def _read(path: str) -> Menu:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb.worksheets[0].iter_rows(values_only=True)
            header = next(rows, None) or ()
            columns = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}
            missing = {"categoria", "producto", "precio"} - columns.keys()
            if missing:
                raise ValueError(f"faltan columnas: {', '.join(sorted(missing))}")

            def cell(row, name):
                i = columns.get(name)
                if i is None or i >= len(row) or row[i] is None:
                    return ""
                return str(row[i]).strip()

            menu: Menu = {}
            for row in rows:
                category = cell(row, "categoria")
                product = cell(row, "producto")
                if not category or not product:
                    continue
                variant = cell(row, "variante") or DEFAULT_VARIANT
                try:
                    price = round(float(cell(row, "precio").replace(",", ".")), 2)
                except ValueError:
                    continue
                menu.setdefault(category, {}).setdefault(product, {})[variant] = price
            return menu
        finally:
            wb.close()
