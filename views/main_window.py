import os
import sys

from PyQt6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QLabel, QPushButton, QListWidget, QComboBox, QTextEdit, QFrame, QDialog, QMessageBox, QLineEdit, QCompleter,
)
from PyQt6.QtGui import QFont, QPixmap, QColor, QPalette, QKeySequence, QShortcut
from PyQt6.QtCore import Qt
from models.menu import MenuData
from models.order import OrderManager
from views.dialogs import (
    AddOrderDialog, EditTableDialog, DeleteItemDialog, PaymentDialog
)
from utils.styles import ThemeManager


class ProvaRestaurant(QMainWindow):
    def __init__(self):
        super().__init__()
        self.menu_data = MenuData()
        self.order_manager = OrderManager()
        self.theme_manager = ThemeManager()

        # 👉 índice para el buscador rápido
        self.quick_index = {}
        self._build_quick_index()

        self.setup_window()
        self.init_ui()
        self.apply_stylesheet()

    def setup_window(self):
        self.setWindowTitle("PROVA - Sistema de Pedidos")
        self.setGeometry(100, 100, 1200, 800)
        self.setMinimumSize(1100, 750)

    def init_ui(self):
        self.main_widget = QWidget()
        self.setCentralWidget(self.main_widget)

        self.main_layout = QHBoxLayout()
        self.main_widget.setLayout(self.main_layout)

        self.setup_left_panel()
        self.setup_right_panel()


    def setup_left_panel(self):
        self.left_panel = QFrame()
        self.left_panel.setObjectName("leftPanel")
        self.left_panel.setMaximumWidth(350)

        layout = QVBoxLayout()
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)
        self.left_panel.setLayout(layout)

        # Logo
        self.setup_logo(layout)

        # Título
        title = QLabel("PROVA")
        title.setFont(QFont("Arial", 28, QFont.Weight.Bold))
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title.setObjectName("appTitle")
        layout.addWidget(title)

        # Subtítulo
        subtitle = QLabel("Sistema de Pedidos")
        subtitle.setFont(QFont("Arial", 12))
        subtitle.setAlignment(Qt.AlignmentFlag.AlignCenter)
        subtitle.setObjectName("appSubtitle")
        layout.addWidget(subtitle)

        # Lista de mesas
        self.table_list = QListWidget()
        self.table_list.setObjectName("tableList")
        self.table_list.setFont(QFont("Arial", 12))
        self.table_list.itemClicked.connect(self.select_table)
        layout.addWidget(QLabel("Mesas/Clientes:"))
        layout.addWidget(self.table_list)

        # Botones de gestión
        self.setup_left_panel_buttons(layout)

        self.main_layout.addWidget(self.left_panel)

    def setup_logo(self, layout):
        self.local_image_label = QLabel()
        self.local_image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.local_image_label.setFixedSize(300, 200)

        # 1) Carpeta donde corre el script o el EXE (PyInstaller)
        if getattr(sys, "frozen", False):
            app_dir = os.path.dirname(sys.executable)
        else:
            # sys.argv[0] apunta al script lanzado; así evitamos rutas relativas a 'assets'
            app_dir = os.path.dirname(os.path.abspath(sys.argv[0]))

        # 2) Buscamos una imagen colocada junto al EXE/py (sin empaquetar)
        #    El usuario debe poner "prova.png" AL LADO del .exe/.py
        candidate_paths = [
            os.path.join(app_dir, "prova.png"),
            os.path.join(app_dir, "PROVA.png"),
        ]

        pixmap = QPixmap()
        for p in candidate_paths:
            if os.path.exists(p):
                pixmap = QPixmap(p)
                break

        if not pixmap.isNull():
            scaled = pixmap.scaled(
                self.local_image_label.size(),
                Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation
            )
            self.local_image_label.setPixmap(scaled)
        else:
            # Fallback si no la ponen
            self.local_image_label.setText("Logo del Local\n(prova.png)")
            self.local_image_label.setStyleSheet("""
                background-color: #F0F0F0;
                border: 2px dashed #AAAAAA;
                color: #666666;
                font-weight: bold;
            """)

        layout.addWidget(self.local_image_label)

    def setup_left_panel_buttons(self, layout):
        btn_layout = QGridLayout()
        btn_layout.setSpacing(10)

        self.add_table_btn = self.create_fancy_button("Nuevo Pedido", "primary", self.add_pedido)
        self.edit_table_btn = self.create_fancy_button("Editar", "secondary", self.edit_table_name)
        self.delete_table_btn = self.create_fancy_button("Eliminar", "danger", self.delete_table)
        self.theme_toggle_btn = self.create_fancy_button("Cambiar Tema", "accent", self.toggle_theme)

        btn_layout.addWidget(self.add_table_btn, 0, 0)
        btn_layout.addWidget(self.edit_table_btn, 0, 1)
        btn_layout.addWidget(self.delete_table_btn, 1, 0)
        btn_layout.addWidget(self.theme_toggle_btn, 1, 1)

        layout.addLayout(btn_layout)

    def setup_right_panel(self):
        self.right_panel = QFrame()
        self.right_panel.setObjectName("rightPanel")

        layout = QVBoxLayout()
        layout.setContentsMargins(30, 30, 30, 30)
        layout.setSpacing(20)
        self.right_panel.setLayout(layout)

        # Encabezado
        self.current_order_header = QLabel("Pedido Actual")
        self.current_order_header.setFont(QFont("Arial", 18, QFont.Weight.Bold))
        self.current_order_header.setObjectName("orderHeader")
        layout.addWidget(self.current_order_header)

        # Info mesa
        self.table_info = QLabel("Seleccione una mesa o cree un nuevo pedido")
        self.table_info.setFont(QFont("Arial", 12))
        self.table_info.setObjectName("tableInfo")
        layout.addWidget(self.table_info)

        # Selectores de platillos
        self.setup_menu_selectors(layout)

        # 👉 aquí insertas el buscador
        self.setup_quick_search(layout)

        # Botones de acción
        self.setup_action_buttons(layout)

        # Display del pedido
        self.order_display = QTextEdit()
        self.order_display.setObjectName("orderDisplay")
        self.order_display.setFont(QFont("Arial", 12))
        self.order_display.setReadOnly(True)
        layout.addWidget(self.order_display)

        # Botón de exportación
        self.export_btn = self.create_fancy_button("Exportar a Excel", "accent", self.save_to_excel)
        layout.addWidget(self.export_btn)

        self.main_layout.addWidget(self.right_panel)

        # Crear widgets para los pagos mixtos
        self.split_payment_label = QLabel("Pago Mixto")
        self.split_payment_label.setFont(QFont("Arial", 12))
        self.cash_amount_input = QLineEdit()
        self.cash_amount_input.setPlaceholderText("Monto en Efectivo")
        self.qr_amount_input = QLineEdit()
        self.qr_amount_input.setPlaceholderText("Monto en QR")

        self.payment_type_combo = QComboBox()
        self.payment_type_combo.addItems(["Efectivo", "QR", "Pago Mixto"])
        self.payment_type_combo.currentIndexChanged.connect(self.toggle_split_payment)

        # Añadir al layout de pagos
        self.payment_layout.addWidget(self.payment_type_combo)
        self.payment_layout.addWidget(self.split_payment_label)
        self.payment_layout.addWidget(self.cash_amount_input)
        self.payment_layout.addWidget(self.qr_amount_input)

        self.change_method_label = QLabel("Cambio entregado en:")
        self.change_method_combo = QComboBox()
        self.change_method_combo.addItems(["Efectivo", "QR"])

        self.payment_layout.addWidget(self.change_method_label)
        self.payment_layout.addWidget(self.change_method_combo)

    def toggle_split_payment(self, index):
        if self.payment_type_combo.currentText() == "Pago Mixto":
            self.split_payment_label.show()
            self.cash_amount_input.show()
            self.qr_amount_input.show()
        else:
            self.split_payment_label.hide()
            self.cash_amount_input.hide()
            self.qr_amount_input.hide()

        self.split_payment_label.hide()
        self.cash_amount_input.hide()
        self.qr_amount_input.hide()

    def setup_menu_selectors(self, layout):
        form_layout = QGridLayout()
        form_layout.setSpacing(15)

        # Categoría
        lbl_category = QLabel("Categoría:")
        lbl_category.setFont(QFont("Arial", 13))
        self.category_combo = QComboBox()
        self.category_combo.addItems(self.menu_data.get_menu_prices().keys())
        self.category_combo.setObjectName("categoryCombo")
        self.category_combo.currentIndexChanged.connect(self.update_dishes)
        form_layout.addWidget(lbl_category, 0, 0)
        form_layout.addWidget(self.category_combo, 0, 1)

        # Platillo
        lbl_dish = QLabel("Platillo:")
        lbl_dish.setFont(QFont("Arial", 13))
        self.dish_combo = QComboBox()
        self.dish_combo.setObjectName("dishCombo")
        self.dish_combo.currentIndexChanged.connect(self.update_variants)
        form_layout.addWidget(lbl_dish, 1, 0)
        form_layout.addWidget(self.dish_combo, 1, 1)

        # Variante
        lbl_variant = QLabel("Variante:")
        lbl_variant.setFont(QFont("Arial", 13))
        self.variant_combo = QComboBox()
        self.variant_combo.setObjectName("variantCombo")
        form_layout.addWidget(lbl_variant, 2, 0)
        form_layout.addWidget(self.variant_combo, 2, 1)

        # Tipo de pedido
        lbl_order_type = QLabel("Tipo:")
        lbl_order_type.setFont(QFont("Arial", 13))
        self.order_type_combo = QComboBox()
        self.order_type_combo.addItems(["Mesa", "Para llevar"])
        self.order_type_combo.setObjectName("orderTypeCombo")
        self.order_type_combo.currentTextChanged.connect(self.set_order_type)
        form_layout.addWidget(lbl_order_type, 3, 0)
        form_layout.addWidget(self.order_type_combo, 3, 1)

        layout.addLayout(form_layout)
        self.update_dishes()
        self.update_variants()

    def toggle_delivery_fields(self, index):
        if self.order_type_combo.currentText() == "Para Llevar":
            self.moto_cost_label.show()
            self.moto_cost_input.show()
            self.moto_payment_method_label.show()
            self.moto_payment_method_combo.show()
        else:
            self.moto_cost_label.hide()
            self.moto_cost_input.hide()
            self.moto_payment_method_label.hide()
            self.moto_payment_method_combo.hide()

    def setup_action_buttons(self, layout):
        action_layout = QHBoxLayout()
        action_layout.setSpacing(15)

        self.add_item_btn = self.create_fancy_button("Agregar", "success", self.add_order)
        self.remove_item_btn = self.create_fancy_button("Eliminar", "warning", self.delete_platillo)
        self.pay_btn = self.create_fancy_button("Pagar", "primary", self.mark_as_paid)
        self.print_btn = self.create_fancy_button("Imprimir", "secondary", self.print_order)

        action_layout.addWidget(self.add_item_btn)
        action_layout.addWidget(self.remove_item_btn)
        action_layout.addWidget(self.pay_btn)
        action_layout.addWidget(self.print_btn)

        layout.addLayout(action_layout)

    def create_fancy_button(self, text, color_type, callback):
        btn = QPushButton(text)
        btn.setFont(QFont("Arial", 11, QFont.Weight.Bold))
        btn.setObjectName(f"{color_type}Button")
        btn.clicked.connect(callback)
        return btn

    def apply_stylesheet(self):
        self.setStyleSheet(self.theme_manager.get_stylesheet())
        palette = self.palette()
        theme = self.theme_manager.get_current_theme()
        palette.setColor(QPalette.ColorRole.Window, QColor(theme['background']))
        self.setPalette(palette)

    def toggle_theme(self):
        self.theme_manager.toggle_theme()
        self.apply_stylesheet()

    def update_dishes(self):
        current_category = self.category_combo.currentText()
        self.dish_combo.clear()
        self.dish_combo.addItems(self.menu_data.get_menu_prices()[current_category].keys())

    def update_variants(self):
        current_category = self.category_combo.currentText()
        current_dish = self.dish_combo.currentText()
        self.variant_combo.clear()

        menu = self.menu_data.get_menu_prices()
        if current_category in menu and current_dish in menu[current_category]:
            self.variant_combo.addItems(menu[current_category][current_dish].keys())

    def set_order_type(self, tipo):
        self.order_manager.order_type = tipo

    def add_pedido(self):
        dialog = AddOrderDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            name = dialog.get_order_name()
            ok, result = self.order_manager.create_table(name)
            if not ok:
                # result es sugerencia o mensaje
                QMessageBox.warning(
                    self,
                    "Nombre en uso",
                    f"Ya existe una mesa con ese nombre.\nPrueba con: «{result}»."
                )
                return

            created_name = result
            self.table_list.addItem(created_name)

    def select_table(self, item):
        table_name = item.text()
        self.order_manager.set_current_table(table_name)
        self.table_info.setText(f"Mesa: {table_name}")
        self.update_order_display()

    def edit_table_name(self):
        if not self.order_manager.current_table:
            QMessageBox.warning(self, "Error", "Selecciona un pedido primero")
            return

        dialog = EditTableDialog(self.order_manager.current_table, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            new_name = dialog.get_new_name()
            if not new_name.strip():
                QMessageBox.warning(self, "Nombre vacío", "Ingresa un nombre válido.")
                return

            # Intentar renombrar; si falla, probablemente porque colisiona
            if not self.order_manager.rename_table(self.order_manager.current_table, new_name):
                suggestion = self.order_manager.suggest_name(new_name)
                QMessageBox.warning(
                    self,
                    "Nombre en uso",
                    f"Ya existe «{new_name}».\nPrueba con: «{suggestion}»."
                )
                return

            # Éxito: refrescar UI
            self.table_list.currentItem().setText(new_name)
            self.order_manager.set_current_table(new_name)
            self.table_info.setText(f"Mesa: {new_name}")

    def delete_table(self):
        if not self.order_manager.current_table:
            QMessageBox.warning(self, "Error", "Selecciona un pedido primero")
            return

        reply = QMessageBox.question(
            self, "Confirmar",
            f"¿Eliminar el pedido de {self.order_manager.current_table}?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if reply == QMessageBox.StandardButton.Yes:
            self.order_manager.delete_table(self.order_manager.current_table)
            self.table_list.takeItem(self.table_list.row(self.table_list.currentItem()))
            self.order_manager.current_table = None
            self.table_info.setText("Seleccione una mesa o cree un nuevo pedido")
            self.order_display.clear()

    def add_order(self):
        if not self.order_manager.current_table:
            QMessageBox.warning(self, "Error", "Selecciona o crea un pedido primero")
            return

        category = self.category_combo.currentText()
        dish = self.dish_combo.currentText()
        variant = self.variant_combo.currentText()
        price = self.menu_data.get_menu_prices()[category][dish][variant]

        self.order_manager.add_item_to_order(category, dish, variant, price)
        self.update_order_display()

    def delete_platillo(self):
        if not self.order_manager.current_table or not self.order_manager.table_orders.get(
                self.order_manager.current_table):
            QMessageBox.warning(self, "Error", "No hay platillos para eliminar")
            return

        dialog = DeleteItemDialog(self.order_manager.table_orders[self.order_manager.current_table], self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            index = dialog.get_selected_index()
            if index is not None:
                self.order_manager.remove_item_from_order(self.order_manager.current_table, index)
                self.update_order_display()

    def update_order_display(self):
        if not self.order_manager.current_table:
            self.order_display.clear()
            return

        order_summary, total = self.order_manager.get_order_summary(self.order_manager.current_table)
        paid_status = self.order_manager.get_payment_status(self.order_manager.current_table)

        text = f"=== {self.order_manager.current_table} ===\n"
        text += f"Tipo: {self.order_manager.order_type}\n"
        text += f"Estado: {'PAGADO (' + paid_status[1] + ')' if paid_status[0] else 'PENDIENTE'}\n"
        text += "-" * 40 + "\n"

        for item, count in order_summary.items():
            price_per_unit = next(
                p['price'] for p in self.order_manager.table_orders[self.order_manager.current_table]
                if f"{p['dish']} ({p['variant']})" == item
            )
            subtotal = price_per_unit * count
            text += f"{item} x{count} = Bs{subtotal}\n"

        text += "-" * 40 + f"\nTOTAL: Bs{total}"
        self.order_display.setPlainText(text)

    def mark_as_paid(self):
        if not self.order_manager.current_table:
            QMessageBox.warning(self, "Error", "Selecciona un pedido primero")
            return

        if not self.order_manager.table_orders.get(self.order_manager.current_table):
            QMessageBox.warning(self, "Error", "El pedido está vacío")
            return

        # Mostrar diálogo de pago con cálculo de cambio
        dialog = PaymentDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            # Registrar pago con todos los detalles
            self.order_manager.set_payment_status(
                self.order_manager.current_table,
                paid=True,
                method=dialog.get_payment_method(),
                amount_paid=dialog.get_amount_paid(),
                change=dialog.get_change()
            )

            # Mensaje de confirmación detallado
            _, total = self.order_manager.get_order_summary(self.order_manager.current_table)
            msg = f"✅ Pago registrado exitosamente\n\n"
            msg += f"Mesa: {self.order_manager.current_table}\n"
            msg += f"Total: Bs. {total:.2f}\n"
            msg += f"Método: {dialog.get_payment_method()}\n"

            if dialog.get_payment_method() == "Efectivo":
                msg += f"Recibido: Bs. {dialog.get_amount_paid():.2f}\n"
                msg += f"💰 Cambio: Bs. {dialog.get_change():.2f}"

            QMessageBox.information(self, "Pago Confirmado", msg)
            self.update_order_display()

            # Actualizar color en la lista
            items = self.table_list.findItems(self.order_manager.current_table, Qt.MatchFlag.MatchExactly)
            if items:
                item = items[0]
                theme = self.theme_manager.get_current_theme()
                item.setBackground(QColor(theme['success']))
                item.setForeground(QColor("white"))

    def print_order(self):
        if not self.order_manager.current_table or not self.order_manager.table_orders.get(
                self.order_manager.current_table):
            QMessageBox.warning(self, "Error", "No hay nada para imprimir")
            return

        try:
            from PyQt6.QtPrintSupport import QPrinter, QPrintDialog
            from PyQt6.QtGui import QPainter, QFont
            from PyQt6.QtCore import Qt

            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            dlg = QPrintDialog(printer, self)

            if dlg.exec() == QDialog.DialogCode.Accepted:
                order_summary, total = self.order_manager.get_order_summary(self.order_manager.current_table)

                # Construir texto de la comanda
                order_details = []
                for item, count in order_summary.items():
                    price_per_unit = next(
                        p['price'] for p in self.order_manager.table_orders[self.order_manager.current_table]
                        if f"{p['dish']} ({p['variant']})" == item
                    )
                    subtotal = price_per_unit * count
                    order_details.append(f"{count}x {item} - {subtotal} bs")

                comanda = (
                    f"           PROVA - Comida Mexicana\n"
                    f"           -------------------------\n\n"
                    f"Pedido: {self.order_manager.current_table}\n"
                    f"Tipo: {self.order_manager.order_type}\n"
                    f"{'-' * 40}\n"
                    f"{chr(10).join(order_details)}\n"
                    f"{'-' * 40}\n"
                    f"TOTAL: {total} bs\n\n"
                    f"¡Gracias por su preferencia!"
                )

                painter = QPainter()
                painter.begin(printer)

                font = QFont("Arial", 12)
                painter.setFont(font)

                rect = painter.viewport()
                painter.drawText(rect, Qt.TextFlag.TextWordWrap, comanda)

                painter.end()

                QMessageBox.information(self, "Éxito", "La comanda se ha enviado a imprimir")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Error al imprimir: {str(e)}")

    def _build_quick_index(self):
        """Crea el índice de búsqueda: 'Platillo (Variante) - BsX' -> (cat, dish, variant, price)"""
        self.quick_index.clear()
        menu = self.menu_data.get_menu_prices()
        for category, dishes in menu.items():
            for dish, variants in dishes.items():
                for variant, price in variants.items():
                    label = f"{dish} ({variant}) - Bs{price}"
                    self.quick_index[label] = (category, dish, variant, float(price))

    def setup_quick_search(self, parent_layout):
        row = QHBoxLayout()

        lbl = QLabel("Búsqueda rápida:")
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Ej: taco, birria, horchata, queso...")

        # Autocompletado (contiene, sin mayúsculas)
        self.completer = QCompleter(list(self.quick_index.keys()))
        self.completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        self.completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.search_input.setCompleter(self.completer)

        self.search_add_btn = self.create_fancy_button("Agregar", "success", self.on_quick_add)
        self.search_input.returnPressed.connect(self.on_quick_add)
        self.completer.activated[str].connect(self.on_quick_choose)

        row.addWidget(lbl)
        row.addWidget(self.search_input, 1)
        row.addWidget(self.search_add_btn)
        parent_layout.addLayout(row)

        # Atajo Ctrl+K
        QShortcut(QKeySequence("Ctrl+K"), self, activated=lambda: self.search_input.setFocus())

    def on_quick_choose(self, text: str):
        if text and text in self.quick_index:
            self._add_from_quick_key(text)

    def on_quick_add(self):
        if not self.order_manager.current_table:
            QMessageBox.warning(self, "Sin mesa", "Selecciona o crea un pedido primero.")
            return
        txt = (self.search_input.text() or "").strip()
        if not txt:
            return

        if txt in self.quick_index:
            self._add_from_quick_key(txt)
            return

        # primer candidato que contenga el texto
        for k in self.quick_index.keys():
            if txt.lower() in k.lower():
                self._add_from_quick_key(k)
                return

        QMessageBox.information(self, "Sin resultados", f"No encontré «{txt}».")

    def _add_from_quick_key(self, key: str):
        category, dish, variant, price = self.quick_index[key]
        self.order_manager.add_item_to_order(category, dish, variant, price)
        self.update_order_display()
        self.search_input.clear()

    def save_to_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font

            wb = Workbook()
            ws = wb.active
            ws.title = "Pedidos"

            # Encabezados
            headers = ["Mesa", "Total (Bs)", "Tipo", "Pagado", "Método de Pago"]
            ws.append(headers)

            # Estilo para encabezados
            for col in range(1, len(headers) + 1):
                ws.cell(1, col).font = Font(bold=True)

            # Datos
            for mesa, items in self.order_manager.table_orders.items():
                total = sum(item['price'] for item in items)
                paid_status = self.order_manager.get_payment_status(mesa)

                ws.append([
                    mesa,
                    total,
                    self.order_manager.order_type,
                    "Sí" if paid_status[0] else "No",
                    paid_status[1] if paid_status[0] else "N/A"
                ])

            # Ajustar anchos de columnas
            for column in ws.columns:
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2) * 1.2
                ws.column_dimensions[column[0].column_letter].width = adjusted_width

            # Guardar archivo
            wb.save("pedidos.xlsx")
            QMessageBox.information(self, "Éxito", "Datos exportados a pedidos.xlsx")

        except Exception as e:
            QMessageBox.critical(self, "Error", f"Error al exportar: {str(e)}")